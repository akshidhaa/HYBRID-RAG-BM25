"""
ingestion.py
------------
Stage 1 of the architecture diagram: Document Ingestion.

Handles text extraction + light normalization for .txt, .docx, and .pdf
files, so the rest of the pipeline (chunking -> indexing -> retrieval ->
abstention) never has to know or care what format the source document was.
"""

import re
import os


class UnsupportedFileType(Exception):
    pass


def _clean_text(text: str) -> str:
    """Text Cleaning & Normalization step from the diagram: collapse
    excess whitespace/blank lines left behind by extraction, without
    altering actual content."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def _extract_docx(path: str) -> str:
    import docx
    doc = docx.Document(path)
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    # Tables often carry real content (e.g. financial figures) -- include them.
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _extract_pdf(path: str) -> str:
    from pypdf import PdfReader
    reader = PdfReader(path)
    parts = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            parts.append(page_text)
    return "\n\n".join(parts)


EXTRACTORS = {
    ".txt": _extract_txt,
    ".md": _extract_txt,
    ".docx": _extract_docx,
    ".pdf": _extract_pdf,
}


def extract_text(path: str) -> str:
    """Given a path to a .txt, .md, .docx, or .pdf file, returns cleaned
    plain text ready for chunking."""
    ext = os.path.splitext(path)[1].lower()
    extractor = EXTRACTORS.get(ext)
    if extractor is None:
        raise UnsupportedFileType(
            f"Unsupported file type '{ext}'. Supported: {', '.join(EXTRACTORS)}"
        )
    raw_text = extractor(path)
    return _clean_text(raw_text)


if __name__ == "__main__":
    text = extract_text("sample_data/document.txt")
    print(text[:300])
    print("...")
    print(f"\nTotal chars: {len(text)}")
